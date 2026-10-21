"""Stage 3: build the Square item-library import workbook.

The header row of config/square_template.xlsx (the blank template downloaded from
Square Dashboard > Items > Item library > Actions > Import) is the contract. Logical
columns are resolved against it by name at runtime, so a renamed column fails loudly here
instead of producing a file Square rejects. If the template is missing, a built-in header
list is used and the summary says so.
"""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from .config import Settings
from .naming import description as make_description
from .naming import item_name as make_item_name
from .naming import record_sku, variation_name as make_variation_name
from .records import Record
from .tiering import BulkBucket, bucket_for, load_buckets


class SquareTemplateError(ValueError):
    pass


class ImportValidationError(ValueError):
    pass


# Fallback only. Square changes these; prefer the real template file.
DEFAULT_TEMPLATE_HEADERS = [
    "Token", "Item Name", "Variation Name", "SKU", "Description", "Category", "GTIN",
    "Reporting Category", "SEO Title", "SEO Description", "Permalink",
    "Square Online Item Visibility", "Shipping Enabled", "Self-serve Ordering Enabled",
    "Delivery Enabled", "Pickup Enabled", "Price", "Online Sale Price", "Archived", "Sellable",
    "Stockable", "Skip Detail Screen in POS", "Option Name 1", "Option Value 1", "Item Type",
    "Weight (lb)", "Social Media Link Title", "Social Media Link Description",
    "Current Quantity {loc}", "New Quantity {loc}", "Stock Alert Enabled {loc}",
    "Stock Alert Count {loc}",
]


def _norm(header: Any) -> str:
    s = re.sub(r"\(locked\)", "", str(header or ""), flags=re.IGNORECASE)
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


@dataclass
class ColumnMap:
    headers: list[str]          # positional; "" for empty header cells
    token: str | None
    item_name: str
    variation_name: str
    sku: str
    description: str | None
    category: str | None
    price: str
    item_type: str | None
    new_quantity: str
    current_quantity: str | None
    stock_alert_enabled: str | None
    stock_alert_count: str | None = None
    tax: list[str] = field(default_factory=list)
    source: str = "template"   # template | builtin | export
    header_row: int = 1         # 1-based row of the header in the source workbook
    template_path: Path | None = None

    def index(self, header: str) -> int:
        return self.headers.index(header)


def find_header_row(ws, max_scan: int = 30) -> tuple[int, list[str]]:
    """Square's template has instruction rows above the header; find the real header row.

    Returns (1-based row index, positional header list with trailing empties trimmed).
    """
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        norms = {_norm(c) for c in row if c not in (None, "")}
        if "item name" in norms and "sku" in norms:
            headers = ["" if c is None else str(c) for c in row]
            while headers and headers[-1] == "":
                headers.pop()
            return i, headers
    raise SquareTemplateError(
        f"no header row with both 'Item Name' and 'SKU' found in the first {max_scan} rows")


def _find(headers: list[str], *names: str) -> str | None:
    wanted = {_norm(n) for n in names}
    return next((h for h in headers if _norm(h) in wanted), None)


def _find_location_column(headers: list[str], prefix: str, location_name: str, required: bool
                          ) -> str | None:
    p = _norm(prefix)
    candidates = [h for h in headers if _norm(h).startswith(p)]
    if location_name:
        loc = _norm(location_name)
        narrowed = [h for h in candidates if loc and loc in _norm(h)]
        if narrowed:
            candidates = narrowed
        elif len(candidates) > 1 and required:
            raise SquareTemplateError(
                f"square.location_name is {location_name!r} but no {prefix!r} column mentions it. "
                f"Columns found: {candidates}. Fix square.location_name in config/settings.yaml.")
        # A single unmatched column is used as-is: with inventory_via api it stays blank anyway,
        # and the template may simply predate a location rename.
    if len(candidates) == 1:
        return candidates[0]
    if not candidates:
        if required:
            raise SquareTemplateError(f"template has no {prefix!r} column")
        return None
    raise SquareTemplateError(
        f"template has several {prefix!r} columns ({candidates}); set square.location_name "
        "in config/settings.yaml to pick one.")


def resolve_columns(headers: list[str], location_name: str = "", source: str = "template",
                    header_row: int = 1, template_path: Path | None = None) -> ColumnMap:
    headers = ["" if h is None else str(h) for h in headers]
    while headers and headers[-1] == "":
        headers.pop()
    item = _find(headers, "Item Name")
    var = _find(headers, "Variation Name")
    sku = _find(headers, "SKU")
    price = _find(headers, "Price")
    missing = [n for n, v in (("Item Name", item), ("Variation Name", var), ("SKU", sku),
                              ("Price", price)) if v is None]
    if missing:
        raise SquareTemplateError(
            f"Square template is missing column(s) {missing}. Headers: {headers}")
    return ColumnMap(
        headers=headers,
        token=_find(headers, "Token"),
        item_name=item, variation_name=var, sku=sku,
        description=_find(headers, "Description"),
        category=_find(headers, "Category", "Categories"),
        price=price,
        item_type=_find(headers, "Item Type"),
        new_quantity=_find_location_column(headers, "New Quantity", location_name, required=True),
        current_quantity=_find_location_column(headers, "Current Quantity", location_name, required=False),
        stock_alert_enabled=_find_location_column(headers, "Stock Alert Enabled", location_name,
                                                  required=False),
        stock_alert_count=_find_location_column(headers, "Stock Alert Count", location_name,
                                                required=False),
        tax=[h for h in headers if _norm(h).startswith("tax ")],
        source=source, header_row=header_row, template_path=template_path,
    )


def read_header_row(xlsx_path: Path) -> tuple[int, list[str]]:
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")      # Square's template has no default style
        wb = load_workbook(xlsx_path, read_only=True)
    return find_header_row(wb.worksheets[0])


def load_template_columns(settings: Settings) -> ColumnMap:
    location = str(settings.get("square.location_name", "") or "")
    template = settings.path("square.template")
    if template.exists():
        header_row, headers = read_header_row(template)
        return resolve_columns(headers, location, source="template", header_row=header_row,
                               template_path=template)
    loc = location or "Main Location"
    headers = [h.replace("{loc}", loc) for h in DEFAULT_TEMPLATE_HEADERS]
    return resolve_columns(headers, location, source="builtin")


# --- rows --------------------------------------------------------------------------------


@dataclass
class ImportRow:
    item_name: str
    variation_name: str
    sku: str
    price: float | None         # None = blank Price cell (variable price, set in Square)
    quantity: int
    description: str = ""
    category: str = ""
    token: str = ""
    tier: str = "tracked"
    extra: dict[str, Any] = field(default_factory=dict)   # carried-over export cells


@dataclass
class BuildResult:
    rows: list[ImportRow]
    warnings: list[str]
    skipped: list[str]
    by_tier: dict[str, int]
    by_set: dict[str, int]
    bulk_detail: list[Record]
    merged_skus: int = 0
    split_items: list[str] = field(default_factory=list)


def _bulk_sku(bucket: BulkBucket) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", bucket.name.lower()).strip("-")
    return f"bulk-{slug}-{'f' if bucket.finish == 'foil' else 'nf'}"


def build_rows(records: list[Record], settings: Settings) -> BuildResult:
    include_lang = bool(settings.get("sku.include_language", False))
    category = str(settings.get("square.category", "MTG Singles") or "")
    cap = int(settings.get("square.variation_cap", 250) or 250)
    buckets = load_buckets(settings.get("tier.bulk_buckets") or [])
    blank_prices = str(settings.get("pricing.source", "scryfall")).lower() == "none"

    warnings: list[str] = []
    skipped: list[str] = []
    by_set: dict[str, int] = defaultdict(int)
    tracked: dict[str, ImportRow] = {}
    tracked_set_of: dict[str, str] = {}
    merged = 0
    bulk_qty: dict[str, int] = defaultdict(int)
    bulk_detail: list[Record] = []

    for rec in records:
        if rec.tier == "bulk":
            b = bucket_for(buckets, rec.rarity, rec.finish)
            if b is None:
                skipped.append(f"{rec.name} ({rec.set} #{rec.collector_number}): bulk with no bucket")
                continue
            bulk_qty[_bulk_sku(b)] += rec.count
            bulk_detail.append(rec)
            continue
        if rec.sell_price is None and not blank_prices:
            skipped.append(f"{rec.name} ({rec.set} #{rec.collector_number} {rec.finish} {rec.condition}): "
                           "no price (listed in needs_manual_price.csv)")
            continue
        sku = record_sku(rec, include_lang)
        if sku in tracked:
            tracked[sku].quantity += rec.count
            merged += 1
            continue
        tracked[sku] = ImportRow(
            item_name=make_item_name(rec),
            variation_name=make_variation_name(rec, include_lang),
            sku=sku, price=None if rec.sell_price is None else float(rec.sell_price),
            quantity=rec.count,
            description=make_description(rec), category=category, tier="tracked",
        )
        tracked_set_of[sku] = rec.set_name
        by_set[rec.set_name] += rec.count

    # Variation cap: split oversized items by set, e.g. basic lands.
    per_item: dict[str, list[ImportRow]] = defaultdict(list)
    for row in tracked.values():
        per_item[row.item_name].append(row)
    split_items: list[str] = []
    for name, rows in per_item.items():
        if len(rows) <= cap:
            continue
        split_items.append(name)
        for row in rows:
            row.item_name = f"{name} ({tracked_set_of[row.sku]})"
        per_set: dict[str, int] = defaultdict(int)
        for row in rows:
            per_set[row.item_name] += 1
        over = [k for k, n in per_set.items() if n > cap]
        if over:
            raise ImportValidationError(
                f"{over} still exceed the variation cap of {cap} after splitting by set")

    rows = sorted(tracked.values(), key=lambda r: (r.item_name.lower(), r.variation_name.lower()))

    # Bulk items appended with summed quantities.
    seen_bulk: set[str] = set()
    for b in buckets:
        sku = _bulk_sku(b)
        if sku in seen_bulk:
            continue
        seen_bulk.add(sku)
        qty = bulk_qty.get(sku, 0)
        if qty == 0:
            continue
        rows.append(ImportRow(item_name=b.name, variation_name=b.variation_name, sku=sku,
                              price=b.price, quantity=qty, description="Bulk singles",
                              category=category, tier="bulk"))

    validate_rows(rows, allow_blank_prices=blank_prices)
    by_tier = {"tracked": sum(r.quantity for r in rows if r.tier == "tracked"),
               "bulk": sum(r.quantity for r in rows if r.tier == "bulk")}
    if merged:
        warnings.append(f"{merged} Moxfield rows shared a SKU with another row and were merged "
                        "(quantities summed)")
    return BuildResult(rows=rows, warnings=warnings, skipped=skipped, by_tier=by_tier,
                       by_set=dict(sorted(by_set.items())), bulk_detail=bulk_detail,
                       merged_skus=merged, split_items=split_items)


def restrict_to_items(result: BuildResult, max_items: int | None = None,
                      only_names: list[str] | None = None) -> BuildResult:
    """Keep only some items (all their variations) for a small test import.

    `max_items` keeps the first N distinct item names (rows are sorted by name);
    `only_names` keeps exactly those names (case-insensitive). Bulk rows are dropped."""
    wanted = {n.strip().lower() for n in only_names} if only_names else None
    kept_names: list[str] = []
    rows: list[ImportRow] = []
    for r in result.rows:
        if r.tier != "tracked":
            continue
        if wanted is not None and r.item_name.lower() not in wanted:
            continue
        if r.item_name not in kept_names:
            if max_items is not None and len(kept_names) >= max_items:
                continue
            kept_names.append(r.item_name)
        rows.append(r)
    by_set: dict[str, int] = defaultdict(int)
    for r in rows:
        # The set name is the part of the variation name before " #".
        by_set[r.variation_name.split(" #")[0]] += r.quantity
    warnings = list(result.warnings) + [
        f"test run: limited to {len(kept_names)} items ({len(rows)} variations); "
        f"{len(result.rows) - len(rows)} rows left out"]
    if wanted is not None:
        missing = sorted(wanted - {n.lower() for n in kept_names})
        if missing:
            warnings.append(f"--only-items names not found in the build: {missing}")
    return BuildResult(rows=rows, warnings=warnings, skipped=result.skipped,
                       by_tier={"tracked": sum(r.quantity for r in rows), "bulk": 0},
                       by_set=dict(sorted(by_set.items())), bulk_detail=[],
                       merged_skus=result.merged_skus, split_items=result.split_items)


def validate_rows(rows: list[ImportRow], allow_blank_prices: bool = False) -> None:
    seen: set[str] = set()
    problems: list[str] = []
    for r in rows:
        if r.sku in seen:
            problems.append(f"duplicate SKU {r.sku}")
        seen.add(r.sku)
        if r.price is None:
            if not allow_blank_prices:
                problems.append(f"{r.sku}: empty price")
        elif r.price <= 0:
            problems.append(f"{r.sku}: non-positive price")
        if not r.item_name or not r.variation_name:
            problems.append(f"{r.sku}: empty item or variation name")
    if problems:
        raise ImportValidationError("; ".join(problems[:20]) + (" ..." if len(problems) > 20 else ""))


# --- update runs: token matching against a Square export --------------------------------


@dataclass
class ExportMatch:
    tokens_applied: int
    orphans: list[dict[str, Any]]
    zeroed: int
    quantity_mode: str = "set"
    added_to_existing: int = 0      # rows whose quantity = Square current + Moxfield count


def read_square_export(path: Path, location_name: str = "") -> tuple[ColumnMap, list[dict[str, Any]]]:
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    header_row, headers = find_header_row(ws)
    cols = resolve_columns(headers, location_name, source="export", header_row=header_row)
    if cols.token is None:
        raise SquareTemplateError("Square export has no Token column; cannot match for update")
    out: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        row = {h: v for h, v in zip(headers, values) if h}
        if any(v not in (None, "") for v in row.values()):
            out.append(row)
    return cols, out


def _to_int(value: Any) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def apply_square_export(rows: list[ImportRow], export_path: Path, location_name: str,
                        zero_missing: bool, quantity_mode: str = "set") -> ExportMatch:
    """Match our rows to a Square library export by SKU.

    Tokens are copied so Square updates in place. quantity_mode "set" writes the Moxfield
    count as the new stock; "add" writes Square's current quantity plus the Moxfield count
    (Moxfield as an intake log that is wiped after each import).
    """
    if quantity_mode not in ("set", "add"):
        raise ValueError(f"quantity_mode must be set or add, got {quantity_mode!r}")
    if zero_missing and quantity_mode == "add":
        raise ValueError("--zero-missing makes no sense in add mode: SKUs absent from an intake "
                         "file are simply unchanged")
    cols, export_rows = read_square_export(export_path, location_name)
    if quantity_mode == "add" and cols.current_quantity is None:
        raise SquareTemplateError("Square export has no 'Current Quantity' column; add mode needs it")
    by_sku = {str(r.get(cols.sku)).strip(): r for r in export_rows if r.get(cols.sku) not in (None, "")}
    ours = {r.sku: r for r in rows}

    applied = 0
    added = 0
    for sku, row in ours.items():
        ex = by_sku.get(sku)
        if ex is None:
            continue
        if ex.get(cols.token):
            row.token = str(ex[cols.token])
            applied += 1
        if quantity_mode == "add":
            row.quantity += _to_int(ex.get(cols.current_quantity))
            added += 1

    orphans: list[dict[str, Any]] = []
    zeroed = 0
    for sku, ex in by_sku.items():
        if sku in ours:
            continue
        orphans.append({"sku": sku, "token": ex.get(cols.token), "item_name": ex.get(cols.item_name),
                        "variation_name": ex.get(cols.variation_name),
                        "current_quantity": ex.get(cols.current_quantity) if cols.current_quantity else "",
                        "price": ex.get(cols.price)})
        if zero_missing:
            price = ex.get(cols.price)
            try:
                price_f = float(price)
            except (TypeError, ValueError):
                price_f = 0.0
            rows.append(ImportRow(
                item_name=str(ex.get(cols.item_name) or ""), variation_name=str(ex.get(cols.variation_name) or ""),
                sku=sku, price=price_f, quantity=0, token=str(ex.get(cols.token) or ""),
                description=str(ex.get(cols.description) or "") if cols.description else "",
                category=str(ex.get(cols.category) or "") if cols.category else "",
                tier="orphan", extra={k: v for k, v in ex.items()},
            ))
            zeroed += 1
    return ExportMatch(tokens_applied=applied, orphans=orphans, zeroed=zeroed,
                       quantity_mode=quantity_mode, added_to_existing=added)


# --- batch guard --------------------------------------------------------------------------

BATCHES_STATE_FILENAME = "import_batches.json"


def batch_fingerprint(rows: list[ImportRow]) -> str:
    import hashlib
    key = "\n".join(f"{r.sku}\t{r.quantity}" for r in sorted(rows, key=lambda r: r.sku))
    return hashlib.sha256(key.encode("utf-8")).hexdigest()[:16]


def record_batch(state_dir: Path, rows: list[ImportRow], output_name: str) -> str | None:
    """Remember this batch; return the earlier build's date if the same batch was built before.

    In add mode a Moxfield export that was already imported and not wiped would double
    stock, so an identical SKU/quantity set is worth a loud warning."""
    import json
    from datetime import datetime
    path = state_dir / BATCHES_STATE_FILENAME
    batches: list[dict[str, Any]] = []
    if path.exists():
        batches = json.loads(path.read_text(encoding="utf-8")).get("batches", [])
    fp = batch_fingerprint(rows)
    earlier = next((b for b in batches if b["fingerprint"] == fp), None)
    batches.append({"fingerprint": fp, "built_at": datetime.now().isoformat(timespec="seconds"),
                    "file": output_name, "rows": len(rows), "quantity": sum(r.quantity for r in rows)})
    state_dir.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"batches": batches[-200:]}, indent=1), encoding="utf-8")
    return f"{earlier['built_at']} ({earlier['file']})" if earlier else None


def write_orphans_csv(path: Path, orphans: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["sku", "token", "item_name", "variation_name",
                                           "current_quantity", "price"])
        w.writeheader()
        w.writerows(orphans)


# --- output files -------------------------------------------------------------------------


def import_row_cells(cols: ColumnMap, r: ImportRow, taxable: bool | None,
                     stock_alert_count: int | None, write_quantities: bool = True) -> list[Any]:
    """One positional list of cell values for an ImportRow, aligned to cols.headers.

    write_quantities=False leaves the quantity and stock-alert columns blank because the
    Inventory API (`mtgsync inventory`) owns stock; the CSV then only creates/updates items."""
    idx = {h: i for i, h in enumerate(cols.headers) if h}
    cells: list[Any] = [None] * len(cols.headers)
    # Carry over any export cells first (zeroed orphans), then overwrite with ours.
    for h, v in r.extra.items():
        if h in idx:
            cells[idx[h]] = v
    cells[idx[cols.item_name]] = r.item_name
    cells[idx[cols.variation_name]] = r.variation_name
    cells[idx[cols.sku]] = r.sku
    cells[idx[cols.price]] = None if r.price is None else round(r.price, 2)
    if write_quantities:
        cells[idx[cols.new_quantity]] = r.quantity
    if cols.token:
        cells[idx[cols.token]] = r.token or None
    if cols.description:
        cells[idx[cols.description]] = r.description or None
    if cols.category:
        cells[idx[cols.category]] = r.category or None
    if cols.item_type:
        cells[idx[cols.item_type]] = None          # blank = Standard
    if not r.extra:
        if write_quantities and cols.stock_alert_enabled:
            cells[idx[cols.stock_alert_enabled]] = "N" if stock_alert_count is None else "Y"
        if write_quantities and cols.stock_alert_count and stock_alert_count is not None:
            cells[idx[cols.stock_alert_count]] = int(stock_alert_count)
        if taxable is not None:
            for h in cols.tax:
                cells[idx[h]] = "Y" if taxable else "N"
    return cells


def write_import_csv(path: Path, cols: ColumnMap, rows: list[ImportRow],
                     taxable: bool | None = None, stock_alert_count: int | None = None,
                     write_quantities: bool = True) -> None:
    """Same content as the xlsx, as UTF-8 CSV with the header in the first line.

    Upload it directly; opening and re-saving it in a spreadsheet tool can mangle SKUs."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(cols.headers)
        for r in rows:
            w.writerow(["" if v is None else v
                        for v in import_row_cells(cols, r, taxable, stock_alert_count, write_quantities)])


def write_import_xlsx(path: Path, cols: ColumnMap, rows: list[ImportRow],
                      taxable: bool | None = None, stock_alert_count: int | None = None,
                      keep_template_rows: bool = False, write_quantities: bool = True) -> None:
    """Write the import workbook.

    Default layout: a single sheet with the template's header row in row 1 and data from
    row 2, which is the layout of Square's own library export. With keep_template_rows the
    template workbook itself is the base (instruction rows, example row, extra sheets kept)
    and data rows go under its header row; empty boilerplate cells are normalised so no
    content-less inline-string cells are written, which Square's importer rejects.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    if keep_template_rows and cols.template_path is not None and cols.template_path.exists():
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = load_workbook(cols.template_path)
        ws = wb.worksheets[0]
        if ws.max_row > cols.header_row:
            ws.delete_rows(cols.header_row + 1, ws.max_row - cols.header_row)
        from openpyxl.cell.cell import MergedCell
        for row in ws.iter_rows(min_row=1, max_row=cols.header_row):
            for cell in row:
                if not isinstance(cell, MergedCell) and cell.value in (None, ""):
                    cell.value = None       # plain empty cell, not a bare inlineStr
        first_data_row = cols.header_row + 1
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Items"
        for j, h in enumerate(cols.headers, start=1):
            ws.cell(row=1, column=j, value=h or None)
        first_data_row = 2

    idx = {h: i for i, h in enumerate(cols.headers) if h}
    for n, r in enumerate(rows):
        cells = import_row_cells(cols, r, taxable, stock_alert_count, write_quantities)
        row_no = first_data_row + n
        for j, v in enumerate(cells, start=1):
            if v is not None:
                ws.cell(row=row_no, column=j, value=v)
        # Force SKUs to text so nothing like `10e-1-nf-nm` is ever read as a number.
        ws.cell(row=row_no, column=idx[cols.sku] + 1).number_format = "@"
    wb.save(path)


def write_bulk_detail_csv(path: Path, records: list[Record]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["name", "set", "set_name", "collector_number", "rarity", "finish", "condition",
                    "lang", "count", "base_price", "bulk_bucket", "moxfield_line"])
        for r in records:
            w.writerow([r.name, r.set, r.set_name, r.collector_number, r.rarity, r.finish,
                        r.condition, r.lang, r.count,
                        "" if r.base_price is None else r.base_price, r.bulk_bucket, r.moxfield_line])


def write_summary_md(path: Path, *, result: BuildResult, cols: ColumnMap, xlsx_path: Path,
                     export_match: ExportMatch | None, extra_sections: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = ["# Square import summary", "", f"File: `{xlsx_path}`", ""]
    if cols.source == "builtin":
        lines += ["> **Warning:** the Square template file (square.template in settings.yaml) was "
                  "not found, so a built-in header list was used. Download the blank template from "
                  "Square Dashboard (Items > Item library > Actions > Import) before importing "
                  "for real.", ""]
    else:
        lines += [f"Template: `{cols.template_path}` (header on row {cols.header_row}; "
                  f"quantity column `{cols.new_quantity}`)", ""]
    lines += ["## Rows", "", "| Tier | Rows | Quantity |", "| --- | --- | --- |"]
    for tier in ("tracked", "bulk", "orphan"):
        n = sum(1 for r in result.rows if r.tier == tier)
        if n:
            q = sum(r.quantity for r in result.rows if r.tier == tier)
            lines.append(f"| {tier} | {n} | {q} |")
    lines.append(f"| **total** | {len(result.rows)} | {sum(r.quantity for r in result.rows)} |")
    lines += ["", "## Tracked quantity by set", "", "| Set | Quantity |", "| --- | --- |"]
    lines += [f"| {s} | {q} |" for s, q in result.by_set.items()]
    if export_match is not None:
        lines += ["", "## Update run", "",
                  f"- Quantity mode: **{export_match.quantity_mode}** "
                  + ("(New Quantity = Square current + Moxfield count)" if export_match.quantity_mode == "add"
                     else "(New Quantity = Moxfield count)"),
                  f"- Tokens matched from Square export: {export_match.tokens_applied}",
                  f"- Rows added onto existing Square stock: {export_match.added_to_existing}",
                  f"- SKUs in Square but not in this file (orphaned_skus.csv): {len(export_match.orphans)}",
                  f"- Orphans zeroed with --zero-missing: {export_match.zeroed}"]
    if result.split_items:
        lines += ["", "## Items split by set (variation cap)", ""] + [f"- {n}" for n in result.split_items]
    if result.warnings or result.skipped:
        lines += ["", "## Warnings", ""]
        lines += [f"- {w}" for w in result.warnings]
        lines += [f"- skipped: {s}" for s in result.skipped[:50]]
        if len(result.skipped) > 50:
            lines.append(f"- ... and {len(result.skipped) - 50} more skipped rows")
    for section in extra_sections or []:
        lines += ["", section]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")

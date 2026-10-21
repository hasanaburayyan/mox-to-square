from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from mtgsync.config import Settings
from mtgsync.enrich import enrich_rows
from mtgsync.moxfield import read_moxfield_csv
from mtgsync.price_stage import price_records
from mtgsync.records import Record
from mtgsync.square_build import (DEFAULT_TEMPLATE_HEADERS, ImportValidationError, ImportRow,
                                  SquareTemplateError, apply_square_export, build_rows,
                                  resolve_columns, validate_rows, write_import_xlsx)


def _priced_records(settings: Settings, moxfield_csv: Path, scryfall_index) -> list[Record]:
    rows, _, _ = read_moxfield_csv(moxfield_csv)
    records, _, _ = enrich_rows(rows, scryfall_index, settings)
    price_records(records, settings)
    return records


def test_builtin_header_resolution():
    headers = [h.replace("{loc}", "Main Store") for h in DEFAULT_TEMPLATE_HEADERS]
    cols = resolve_columns(headers, "Main Store", source="builtin")
    assert cols.new_quantity == "New Quantity Main Store"
    assert cols.stock_alert_enabled == "Stock Alert Enabled Main Store"
    assert cols.token == "Token" and cols.sku == "SKU"


def test_location_name_must_match_a_quantity_column():
    headers = ["Token", "Item Name", "Variation Name", "SKU", "Price",
               "New Quantity Store A", "New Quantity Store B"]
    with pytest.raises(SquareTemplateError, match="several"):
        resolve_columns(headers, "")
    assert resolve_columns(headers, "Store B").new_quantity == "New Quantity Store B"
    with pytest.raises(SquareTemplateError, match="location_name"):
        resolve_columns(headers, "Store C")
    # A lone quantity column is accepted even if the name differs (template predates a rename).
    single = ["Item Name", "Variation Name", "SKU", "Price", "New Quantity Old Name"]
    assert resolve_columns(single, "New Name").new_quantity == "New Quantity Old Name"


def test_missing_sku_column_is_a_contract_failure():
    with pytest.raises(SquareTemplateError, match="SKU"):
        resolve_columns(["Item Name", "Variation Name", "Price", "New Quantity X"], "")


def test_build_rows_model_a(settings: Settings, moxfield_csv: Path, scryfall_index):
    settings.set("sku.include_language", False)
    records = _priced_records(settings, moxfield_csv, scryfall_index)
    result = build_rows(records, settings)
    by_sku = {r.sku: r for r in result.rows}

    # Two English rows (different tags) plus the Japanese and German rows merge into one SKU
    # (4 + 2 + 1 + 1) because sku.include_language is off.
    assert by_sku["sta-42-nf-nm"].quantity == 8
    assert result.merged_skus == 3
    assert by_sku["sta-42-nf-nm"].price == 2.50          # 2.40 rounded up to 0.25
    assert by_sku["sta-42-nf-nm"].variation_name == "Strixhaven Mystical Archive #42 · NM · Nonfoil"
    assert by_sku["sta-42-e-lp"].price == 4.00           # etched 3.90 -> 4.00
    # Japanese printing shares the SKU unless sku.include_language is on; merged quantity.
    assert "sta-42-nf-nm-ja" not in by_sku
    assert by_sku["neo-141-nf-nm"].item_name == "Fable of the Mirror-Breaker // Reflection of Kiki-Jiki"
    assert by_sku["pdom-207star-f-nm"].price == 31.00
    # Null Scryfall price -> bulk bucket price and needs_manual_price flag.
    bears = next(r for r in records if r.name == "Grizzly Bears" and r.finish == "nonfoil")
    assert bears.needs_manual_price and by_sku["4ed-246-nf-nm"].price == 0.25
    # Model A: no bulk rows.
    assert not [r for r in result.rows if r.tier == "bulk"]
    # Proxy dropped, unmatched promo not present.
    assert not any("lea" in s for s in by_sku)
    assert all(r.category == "MTG Singles" for r in result.rows)


def test_build_rows_language_suffix(settings: Settings, moxfield_csv: Path, scryfall_index):
    settings.set("sku.include_language", True)
    records = _priced_records(settings, moxfield_csv, scryfall_index)
    result = build_rows(records, settings)
    by_sku = {r.sku: r for r in result.rows}
    assert by_sku["sta-42-nf-nm"].quantity == 6            # only the two English rows merge
    assert by_sku["sta-42-nf-nm-ja"].quantity == 1
    assert by_sku["sta-42-nf-nm-ja"].variation_name == "Strixhaven Mystical Archive #42 · NM · Nonfoil · JA"
    assert by_sku["sta-42-nf-nm-ja"].price == 6.00         # own Scryfall object, own price
    assert by_sku["sta-42-nf-nm-de"].price == 2.50         # English fallback object, English price
    assert result.merged_skus == 1


def test_build_rows_model_c_buckets(settings: Settings, moxfield_csv: Path, scryfall_index):
    settings.set("tier.threshold_usd", 1.00)
    records = _priced_records(settings, moxfield_csv, scryfall_index)
    result = build_rows(records, settings)
    bulk = {r.sku: r for r in result.rows if r.tier == "bulk"}
    # Forests (31) + Opt nonfoil (9) + Grizzly Bears nonfoil (2) -> C/U nonfoil bucket.
    assert bulk["bulk-bulk-common-uncommon-nf"].quantity == 31 + 9 + 2
    assert bulk["bulk-bulk-common-uncommon-nf"].price == 0.25
    assert bulk["bulk-bulk-common-uncommon-f"].quantity == 1 + 1   # Opt foil + Bears foil
    assert len(result.bulk_detail) == 8      # 4 Forest rows + 2 Opt + 2 Grizzly Bears
    tracked = {r.sku for r in result.rows if r.tier == "tracked"}
    assert "sta-42-nf-nm" in tracked and "sld-63a-f-nm" in tracked
    assert not any(s.startswith("eld-") for s in tracked)


def test_variation_cap_splits_by_set(settings: Settings, moxfield_csv: Path, scryfall_index):
    settings.set("sku.include_language", False)
    settings.set("square.variation_cap", 3)   # Forest has 4 variations, Lightning Bolt 3
    records = _priced_records(settings, moxfield_csv, scryfall_index)
    result = build_rows(records, settings)
    forests = [r for r in result.rows if r.item_name.startswith("Forest")]
    assert {r.item_name for r in forests} == {"Forest (The Lost Caverns of Ixalan)",
                                              "Forest (Murders at Karlov Manor)",
                                              "Forest (Wilds of Eldraine)"}
    assert result.split_items == ["Forest"]
    settings.set("square.variation_cap", 1)
    with pytest.raises(ImportValidationError, match="variation cap"):
        build_rows(records, settings)


def test_real_square_template(fixtures_dir: Path, tmp_path: Path):
    """Square's real template: instructions in rows 1-5, header on row 6, '(locked)' suffixes."""
    from mtgsync.square_build import read_header_row

    template = fixtures_dir / "square_template.xlsx"
    header_row, headers = read_header_row(template)
    assert header_row == 6
    cols = resolve_columns(headers, "The Warp Gate", source="template", header_row=header_row,
                           template_path=template)
    assert cols.new_quantity == "New Quantity The Warp Gate"
    assert cols.current_quantity == "Current Quantity The Warp Gate"
    assert cols.category == "Categories"
    assert cols.token == "Token" and cols.item_type == "Item Type"
    assert cols.stock_alert_count == "Stock Alert Count The Warp Gate"
    assert len(cols.tax) == 2

    rows = [ImportRow("Lightning Bolt", "STA #42 · NM · Nonfoil", "sta-42-nf-nm", 2.5, 4,
                      description="Instant", category="MTG Singles")]
    out = tmp_path / "import.xlsx"
    write_import_xlsx(out, cols, rows, taxable=False, stock_alert_count=4)

    # Default layout: header in row 1, data from row 2, nothing else.
    wb = load_workbook(out)
    ws = wb.worksheets[0]
    assert wb.sheetnames == ["Items"]
    assert [c.value for c in ws[1]] == headers
    data = {h: v for h, v in zip(headers, [c.value for c in ws[2]])}
    assert data["SKU"] == "sta-42-nf-nm" and data["Price"] == 2.5
    assert data["New Quantity The Warp Gate"] == 4
    assert data["Categories"] == "MTG Singles"
    assert data["Stock Alert Enabled The Warp Gate"] == "Y"
    assert data["Stock Alert Count The Warp Gate"] == 4
    assert all(data[t] == "N" for t in cols.tax)
    assert data["Item Type"] is None
    assert ws.max_row == 2

    # Inventory via API: quantity and alert cells blank, everything else unchanged.
    write_import_xlsx(out, cols, rows, taxable=False, stock_alert_count=4, write_quantities=False)
    ws = load_workbook(out).worksheets[0]
    data = {h: v for h, v in zip(headers, [c.value for c in ws[2]])}
    assert data["SKU"] == "sta-42-nf-nm" and data["New Quantity The Warp Gate"] is None
    assert data["Stock Alert Enabled The Warp Gate"] is None and data["Stock Alert Count The Warp Gate"] is None

    # Alert off: Enabled = N and the count stays blank.
    write_import_xlsx(out, cols, rows, taxable=None, stock_alert_count=None)
    ws = load_workbook(out).worksheets[0]
    data = {h: v for h, v in zip(headers, [c.value for c in ws[2]])}
    assert data["Stock Alert Enabled The Warp Gate"] == "N"
    assert data["Stock Alert Count The Warp Gate"] is None
    assert all(data[t] is None for t in cols.tax)

    # Template-preserving layout: boilerplate kept, no content-less inlineStr cells.
    import zipfile
    write_import_xlsx(out, cols, rows, keep_template_rows=True)
    wb = load_workbook(out)
    ws = wb.worksheets[0]
    assert "Component Inventory" in wb.sheetnames
    assert str(ws.cell(row=3, column=1).value).startswith("Fill in your catalog")
    assert [c.value for c in ws[6]][:6] == headers[:6]
    assert ws.cell(row=7, column=headers.index("SKU") + 1).value == "sta-42-nf-nm"
    assert ws.max_row == 7
    xml = zipfile.ZipFile(out).read("xl/worksheets/sheet1.xml").decode()
    assert 't="inlineStr" />' not in xml and 't="inlineStr"/>' not in xml


def test_restrict_to_items(settings: Settings, moxfield_csv: Path, scryfall_index):
    from mtgsync.square_build import restrict_to_items

    settings.set("tier.threshold_usd", 1.00)      # so there are bulk rows to drop
    records = _priced_records(settings, moxfield_csv, scryfall_index)
    full = build_rows(records, settings)
    limited = restrict_to_items(full, 2)
    names = [r.item_name for r in limited.rows]
    assert sorted(set(names)) == ["Fable of the Mirror-Breaker // Reflection of Kiki-Jiki", "Island"]
    assert all(r.tier == "tracked" for r in limited.rows)
    assert any("limited to 2 items" in w for w in limited.warnings)
    assert limited.by_tier["bulk"] == 0

    only = restrict_to_items(full, only_names=["lightning bolt", "Nope"])   # Forest is bulk here
    assert {r.item_name for r in only.rows} == {"Lightning Bolt"}
    assert any("not found" in w and "nope" in w for w in only.warnings)


def test_description_constant_within_item(settings: Settings, moxfield_csv: Path, scryfall_index):
    """Square drops an item whose variation rows carry different descriptions."""
    records = _priced_records(settings, moxfield_csv, scryfall_index)
    result = build_rows(records, settings)
    per_item: dict[str, set[str]] = {}
    for r in result.rows:
        per_item.setdefault(r.item_name, set()).add((r.description, r.category))
    assert all(len(v) == 1 for v in per_item.values()), per_item
    assert len({r.variation_name.split(" #")[0] for r in result.rows if r.item_name == "Forest"}) == 3


def test_batch_guard(tmp_path: Path):
    from mtgsync.square_build import record_batch
    rows = [ImportRow("A", "v", "a-1-nf-nm", None, 2), ImportRow("B", "v", "b-1-nf-nm", None, 1)]
    assert record_batch(tmp_path, rows, "first.csv") is None
    earlier = record_batch(tmp_path, list(reversed(rows)), "second.csv")   # order irrelevant
    assert earlier and "first.csv" in earlier
    rows[0].quantity = 3
    assert record_batch(tmp_path, rows, "third.csv") is None


def test_validate_rows_rejects_duplicates_and_empty_prices():
    rows = [ImportRow("A", "v1", "sku-1", 1.0, 1), ImportRow("A", "v2", "sku-1", 1.0, 1)]
    with pytest.raises(ImportValidationError, match="duplicate SKU"):
        validate_rows(rows)
    with pytest.raises(ImportValidationError, match="price"):
        validate_rows([ImportRow("A", "v1", "sku-1", 0.0, 1)])


def _fake_square_export(path: Path, location: str, rows: list[dict]) -> None:
    headers = [h.replace("{loc}", location) for h in DEFAULT_TEMPLATE_HEADERS]
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])
    wb.save(path)


def test_token_matching_and_orphans(tmp_path: Path):
    loc = "Main Store"
    export = tmp_path / "square_export.xlsx"
    _fake_square_export(export, loc, [
        {"Token": "TOK1", "Item Name": "Lightning Bolt", "Variation Name": "v", "SKU": "sta-42-nf-nm",
         "Price": 2.5, f"Current Quantity {loc}": 6},
        {"Token": "TOK2", "Item Name": "Gone Card", "Variation Name": "v", "SKU": "xxx-1-nf-nm",
         "Price": 9.0, f"Current Quantity {loc}": 1},
    ])
    rows = [ImportRow("Lightning Bolt", "v", "sta-42-nf-nm", 2.5, 4),
            ImportRow("New Card", "v", "new-1-nf-nm", 1.0, 1)]

    match = apply_square_export(rows, export, loc, zero_missing=False)
    assert match.tokens_applied == 1 and rows[0].token == "TOK1" and rows[1].token == ""
    assert [o["sku"] for o in match.orphans] == ["xxx-1-nf-nm"]
    assert match.zeroed == 0 and len(rows) == 2

    # add mode: matched SKU gets Square current (6) + Moxfield (4); new SKU keeps its count.
    add_rows = [ImportRow("Lightning Bolt", "v", "sta-42-nf-nm", 2.5, 4),
                ImportRow("New Card", "v", "new-1-nf-nm", 1.0, 1)]
    match = apply_square_export(add_rows, export, loc, zero_missing=False, quantity_mode="add")
    assert match.quantity_mode == "add" and match.added_to_existing == 1
    assert add_rows[0].quantity == 10 and add_rows[0].token == "TOK1"
    assert add_rows[1].quantity == 1 and add_rows[1].token == ""
    with pytest.raises(ValueError, match="zero-missing"):
        apply_square_export(add_rows, export, loc, zero_missing=True, quantity_mode="add")

    match = apply_square_export(rows, export, loc, zero_missing=True)
    assert match.zeroed == 1 and rows[-1].sku == "xxx-1-nf-nm" and rows[-1].quantity == 0
    assert rows[-1].token == "TOK2" and rows[-1].tier == "orphan"

    # Written workbook carries tokens, text SKUs and the zeroed row.
    headers = [h.replace("{loc}", loc) for h in DEFAULT_TEMPLATE_HEADERS]
    cols = resolve_columns(headers, loc, source="builtin")
    out = tmp_path / "import.xlsx"
    write_import_xlsx(out, cols, rows)
    ws = load_workbook(out).active
    got = {r[headers.index("SKU")]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert got["sta-42-nf-nm"][headers.index("Token")] == "TOK1"
    assert got["xxx-1-nf-nm"][headers.index(f"New Quantity {loc}")] == 0
    assert got["new-1-nf-nm"][headers.index("Token")] is None
    assert got["new-1-nf-nm"][headers.index(f"Stock Alert Enabled {loc}")] == "N"

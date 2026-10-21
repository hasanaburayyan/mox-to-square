"""Stages 1-3 through the CLI against the fixtures, fully offline."""

import json
from pathlib import Path

from openpyxl import load_workbook

from mtgsync.cli import main


def test_run_offline(project_root: Path, moxfield_csv: Path, capsys):
    rc = main(["--root", str(project_root), "run", str(moxfield_csv), "--offline",
               "--price-source", "scryfall", "--no-square-export"])
    assert rc == 0, capsys.readouterr().err

    out = project_root / "output"
    unmatched = (out / "unmatched.csv").read_text(encoding="utf-8").splitlines()
    assert len(unmatched) == 2                       # header + pfake promo
    assert any("pfake" in line for line in unmatched)
    # The German Bolt has no own Scryfall object and fell back to the English one.
    enriched = (project_root / "work" / "enriched.jsonl").read_text(encoding="utf-8")
    assert '"lang": "de", "lang_fallback": true' in enriched

    manual = (out / "needs_manual_price.csv").read_text(encoding="utf-8").splitlines()
    assert len(manual) == 3                          # header + Grizzly Bears nonfoil + foil

    tracked = json.loads((project_root / "state" / "tracked.json").read_text())["tracked"]
    assert "sta-42" in tracked and "lea-232" not in tracked

    import csv
    path = next(out.glob("square_import_*.csv"))           # csv is the default format
    with open(path, newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    skus = [r["SKU"] for r in rows]
    assert len(skus) == len(set(skus))
    assert "sta-42-nf-nm" in skus and "neo-141-nf-nm" in skus
    assert all(r["Price"] for r in rows)
    assert all(r["Item Type"] == "" for r in rows)
    assert not list(out.glob("*.xlsx"))

    summary = (out / "import_summary.md").read_text(encoding="utf-8")
    assert "built-in header list" in summary        # no template file in the tmp root
    assert "| tracked |" in summary

    # Second run is idempotent and picks up the persisted tracked state.
    assert main(["--root", str(project_root), "run", str(moxfield_csv), "--offline",
                 "--price-source", "scryfall", "--no-square-export"]) == 0


def test_add_mode_requires_square_export(project_root: Path, moxfield_csv: Path, capsys):
    """square.quantity_mode: add (the config default) refuses to build blind, and warns when
    the identical batch is built twice."""
    rc = main(["--root", str(project_root), "run", str(moxfield_csv), "--offline"])
    assert rc == 2
    assert "add mode needs a fresh Square library export" in capsys.readouterr().err
    assert main(["--root", str(project_root), "build", "--no-square-export"]) == 0
    capsys.readouterr()
    assert main(["--root", str(project_root), "build", "--no-square-export"]) == 0
    assert "this exact batch" in capsys.readouterr().err
    assert main(["--root", str(project_root), "build", "--quantity-mode", "set"]) == 0


def test_run_with_no_price_source(project_root: Path, moxfield_csv: Path, capsys):
    """pricing.source: none (the config default): Price blank, nothing flagged for manual pricing."""
    import csv
    rc = main(["--root", str(project_root), "run", str(moxfield_csv), "--offline", "--no-square-export"])
    assert rc == 0, capsys.readouterr().err
    out = project_root / "output"
    with open(next(out.glob("square_import_*.csv")), newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    # 19 Moxfield rows - 1 proxy - 1 unmatched = 17 records; two share a SKU -> 16 rows, none skipped
    assert len(rows) == 16
    assert all(r["Price"] == "" for r in rows)
    # inventory_via: api (config default): quantities and alerts are not in the CSV.
    qty_col = next(h for h in rows[0] if h.startswith("New Quantity"))
    alert_col = next(h for h in rows[0] if h.startswith("Stock Alert Enabled"))
    assert all(r[qty_col] == "" and r[alert_col] == "" for r in rows)
    manual = (out / "needs_manual_price.csv").read_text(encoding="utf-8").splitlines()
    assert len(manual) == 1                              # header only


def test_build_xlsx_format_still_available(project_root: Path, moxfield_csv: Path):
    assert main(["--root", str(project_root), "run", str(moxfield_csv), "--offline",
                 "--price-source", "scryfall", "--format", "xlsx", "--tag", "t",
                 "--no-square-export"]) == 0
    path = next((project_root / "output").glob("square_import_*_t.xlsx"))
    ws = load_workbook(path).active
    headers = [c.value for c in ws[1]]
    skus = [r[headers.index("SKU")] for r in ws.iter_rows(min_row=2, values_only=True)]
    assert {"sta-42-nf-nm", "neo-141-nf-nm"} <= set(skus)


def test_compare_prices_mode(project_root: Path, moxfield_csv: Path):
    assert main(["--root", str(project_root), "enrich", str(moxfield_csv), "--offline"]) == 0
    assert main(["--root", str(project_root), "price", "--compare-prices"]) == 0
    lines = (project_root / "output" / "price_comparison.csv").read_text().splitlines()
    assert lines[0].startswith("name,set,collector_number")
    assert len(lines) == 1 + 17                      # one per enriched record
    assert not (project_root / "work" / "priced.jsonl").exists()
